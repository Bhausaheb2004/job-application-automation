"""Save tailored resume/cover letter as PDF, DOCX and/or Markdown."""
import re
from xml.sax.saxutils import escape

from docx import Document
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer


# ---------- DOCX ----------
def _add_runs(paragraph, text):
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part)


def md_to_docx(md_text, path):
    doc = Document()
    for line in md_text.splitlines():
        line = line.rstrip()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.lstrip().startswith(("- ", "* ")):
            _add_runs(doc.add_paragraph(style="List Bullet"), line.lstrip()[2:])
        else:
            _add_runs(doc.add_paragraph(), line)
    doc.save(path)


# ---------- PDF ----------
def _inline(text):
    """Escape XML and convert **bold** to <b>."""
    text = escape(text)
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)


def md_to_pdf(md_text, path):
    base = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=base["Normal"], fontName="Helvetica",
                          fontSize=10, leading=13, alignment=TA_LEFT)
    h1 = ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=18, leading=22, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=12, leading=15,
                        spaceBefore=8, spaceAfter=3, textColor="#1f3a5f")
    h3 = ParagraphStyle("h3", parent=body, fontName="Helvetica-Bold", fontSize=10.5, leading=13, spaceBefore=4)

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm)
    story, bullets = [], []

    def flush():
        nonlocal bullets
        if bullets:
            story.append(ListFlowable(
                [ListItem(Paragraph(_inline(b), body), leftIndent=12) for b in bullets],
                bulletType="bullet", start="\u2022", leftIndent=12))
            bullets = []

    for line in md_text.splitlines():
        line = line.rstrip()
        if not line:
            flush()
            story.append(Spacer(1, 3))
        elif line.startswith("### "):
            flush(); story.append(Paragraph(_inline(line[4:]), h3))
        elif line.startswith("## "):
            flush(); story.append(Paragraph(_inline(line[3:]), h2))
        elif line.startswith("# "):
            flush(); story.append(Paragraph(_inline(line[2:]), h1))
        elif line.lstrip().startswith(("- ", "* ")):
            bullets.append(line.lstrip()[2:])
        else:
            flush(); story.append(Paragraph(_inline(line), body))
    flush()
    doc.build(story)


# ---------- Save ----------
def save_outputs(folder, resume_md, cover_md, formats=("pdf", "docx", "md")):
    items = [("resume", resume_md)]
    if cover_md:
        items.append(("cover_letter", cover_md))
    for name, text in items:
        if "md" in formats:
            (folder / f"{name}.md").write_text(text, encoding="utf-8")
        if "docx" in formats:
            md_to_docx(text, folder / f"{name}.docx")
        if "pdf" in formats:
            md_to_pdf(text, folder / f"{name}.pdf")
